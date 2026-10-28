"""Сборка отчётов по зарплатам: картинка (PNG) и Excel на каждый отчёт.

Запуск: python build_reports.py <data.json> <папка_вывода> [id отчёта ...]
data.json лежит вне git (.cache/salary-reports/ГГГГ-ММ/), в нём ФИО и ЗП.
"""
import json
import re
import sys
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Шкала выводов (раздел 7 SKILL.md)
ST = {
    "low": ("Ниже рынка", "#F8D7DA"),
    "ok": ("В рынке", "#D4EDDA"),
    "bit": ("Чуть выше", "#FFF3CD"),
    "high": ("Выше рынка", "#D6E4F5"),
}


def money(v):
    return f"{int(round(v)):,}".replace(",", " ") + " ₽"


def pct(d):
    return f"{d:+.1f}%".replace(".", ",").replace("-", "−")


def status(d):
    return "low" if d < -7 else "ok" if d <= 10 else "bit" if d <= 15 else "high"


def conclusion(d, gap):
    """Нейтральный вывод: только факты рынка, без рекомендаций."""
    if d < -15:
        return f"Ниже рынка на {money(gap)}. Есть вероятность, что сотрудник начнёт присматриваться к вакансиям"
    if d < -7:
        return f"Ниже рынка на {money(gap)}"
    if d <= 10:
        return "В рынке"
    if d <= 15:
        return "Немного выше рынка"
    return "Выше рынка"


def nbsp_numbers(s):
    return re.sub(r"(\d) (\d)", "\\1\u00a0\\2", s)


def build(rep, data, out):
    directors = rep["kind"] == "directors"
    rows = []
    for r in rep["rows"]:
        d = round(r["zp"] / r["market"] * 100 - 100, 2)
        rec = r.get("conclusion") or conclusion(d, r["market"] - r["zp"])
        if r.get("note"):
            rec += f" ({r['note']})"
        if directors and r.get("premium") and not r.get("conclusion"):
            rec += ". С премией ~12 500 ₽/мес (оценка)"
        mk = money(r["market"]) + (f" ({r['market_region']})" if r.get("market_region") else "")
        rows.append(dict(r, d=d, st=status(d), rec=rec, mk=mk))
    rows.sort(key=lambda x: x["d"])

    region_col = any(r.get("market_region") for r in rows)
    H = ["№", "Сотрудник", "ЗП сотрудника", "Рынок" if region_col else "Рынок по ЧР",
         "Отклонение", "Статус", "Вывод"] + (["Комментарий"] if directors else [])

    # картинка
    cells = []
    for i, r in enumerate(rows):
        c = [str(i + 1), r["name"], money(r["zp"]), r["mk"], pct(r["d"]), ST[r["st"]][0],
             "\n".join(textwrap.wrap(r["rec"], 31 if directors else 46))]
        if directors:
            c.append("\n".join(textwrap.wrap(nbsp_numbers(r.get("comment") or "Нет данных"), 64)))
        cells.append(c)
    nl = [max(1, *[x.count("\n") + 1 for x in c[6:]]) for c in cells]
    units = [1.3] + nl
    tot = sum(units)
    fig, ax = plt.subplots(figsize=(20 if directors else 14, 0.3 * tot))
    ax.axis("off")
    ax.set_title(f"{rep['heading']} — мониторинг зарплат, {data['month']}  ·  до НДФЛ  ·  средняя за {data['period']}",
                 fontsize=12, fontweight="bold", loc="left", pad=8)
    widths = ([0.02, 0.175, 0.075, 0.07, 0.065, 0.065, 0.165, 0.365] if directors
              else [0.03, 0.2, 0.11, 0.13, 0.095, 0.095, 0.34])
    t = ax.table(cellText=cells, colLabels=H, cellLoc="left", colWidths=widths, bbox=[0, 0, 1, 1])
    t.auto_set_font_size(False)
    t.set_fontsize(9.5)
    for (ri, ci), cell in t.get_celld().items():
        cell.set_edgecolor("#bbbbbb")
        cell.PAD = 0.03
        cell.set_height(units[ri] / tot)
        if ri == 0:
            cell.set_facecolor("#29619b")
            cell.get_text().set_color("white")
            cell.get_text().set_fontweight("bold")
            continue
        if ci in (4, 5):
            cell.set_facecolor(ST[rows[ri - 1]["st"]][1])
        if ci in (2, 3, 4):
            cell.get_text().set_horizontalalignment("right")
            cell._loc = "right"
    plt.savefig(f"{out}/{rep['id']}.png", dpi=150, bbox_inches="tight")
    plt.close(fig)

    # Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Отдел"
    ws.append([f"{rep['title']} — мониторинг зарплат, {data['month']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Суммы до НДФЛ · ЗП — средняя за {data['period']} · рынок — середина вакансий HH"])
    ws.append([])
    ws.append(H)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="29619B")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BBBBBB")
    for i, r in enumerate(rows):
        ws.append([i + 1, r["name"], r["zp"], r["market"], r["d"] / 100, ST[r["st"]][0], r["rec"]]
                  + ([r.get("comment") or "Нет данных"] if directors else []))
        n = ws.max_row
        for col in "CD":
            ws[f"{col}{n}"].number_format = '# ##0 "₽"'
        ws[f"E{n}"].number_format = "+0.0%;-0.0%;0.0%"
        for col in "EF":
            ws[f"{col}{n}"].fill = PatternFill("solid", fgColor=ST[r["st"]][1][1:])
        for c in ws[n]:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, wd in zip("ABCDEFGH", [5, 28, 15, 14, 12, 13, 40 if directors else 55, 80]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"
    low = [f"{r['name']} ({pct(r['d'])})" for r in rows if r["st"] == "low"]
    hi = [f"{r['name']} ({pct(r['d'])})" for r in rows if r["st"] in ("bit", "high")]
    ws.append([])
    ws.append(["Итог: кто не в рынке"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([f"Ниже рынка ({len(low)}): " + ", ".join(low)])
    if hi:
        ws.append([f"Выше рынка ({len(hi)}): " + ", ".join(hi)])
    wb.save(f"{out}/{rep['id']}.xlsx")
    print(rep["id"], "→ ok,", len(rows), "чел.")


def intro(rep, data):
    """Вступление к письму (согласовано 28.10.2026)."""
    if rep["kind"] == "directors":
        return ("Добрый день! Закончила полугодовой мониторинг зарплат, делюсь результатами по руководителям "
                "и сотрудникам в прямом подчинении.\n\n[b]Коротко, как считала[/b]\nСравнила зарплаты с вакансиями "
                f"конкурентов на HH ({rep['region']}). Брала середину рынка, все суммы до НДФЛ. Наши цифры — средняя "
                f"за {data['period']}; у руководителей и получателей премии добавлена полугодовая премия оценочно 12 500 ₽/мес.")
    return (f"Привет, {rep['greeting']}! Закончила полугодовой мониторинг зарплат, делюсь результатами по твоему отделу.\n\n"
            "[b]Коротко, как считала[/b]\nСравнила зарплаты сотрудников отдела с вакансиями конкурентов на HH "
            f"({rep['region']}). Брала середину рынка, все суммы до НДФЛ. Наши цифры — средняя за {data['period']}.")


if __name__ == "__main__":
    data = json.load(open(sys.argv[1]))
    out = sys.argv[2]
    only = set(sys.argv[3:])
    for rep in data["reports"]:
        if only and rep["id"] not in only:
            continue
        build(rep, data, out)
        json.dump({"intro": intro(rep, data)}, open(f"{out}/{rep['id']}.intro.json", "w"), ensure_ascii=False)
